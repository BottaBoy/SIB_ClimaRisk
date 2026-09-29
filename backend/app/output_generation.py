from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any
import json
import math
import mimetypes

from openpyxl import Workbook

from .job_store import JobStore


UTC = timezone.utc
REPO_ROOT = Path(__file__).resolve().parents[2]
BASIN_OVERLAYS_PATH = REPO_ROOT / "web" / "hazard-maps" / "basin-leaflet-overlays.json"
HAZARDS = ("storm", "storm_cmcc")
SCENARIOS = (
    ("annual", "EAI"),
    ("rp50", "RP50"),
    ("rp100", "RP100"),
    ("rp1000", "RP1000"),
)
STATE_COLORS = {
    "S0": "#6AB96F",
    "S1": "#FFE48A",
    "S2": "#FF9138",
    "S3": "#111111",
}
STATE_THRESHOLDS = (
    ("S3", 0.35),
    ("S2", 0.15),
    ("S1", 0.05),
)


def _num(value: Any) -> float:
    try:
        out = float(value or 0.0)
    except Exception:
        return 0.0
    return out


def _flatten(prefix: str, value: Any, out: dict[str, Any]) -> None:
    if isinstance(value, dict):
        for key, item in value.items():
            child = f"{prefix}.{key}" if prefix else str(key)
            _flatten(child, item, out)
        return
    if isinstance(value, list):
        out[prefix] = json.dumps(value, ensure_ascii=False)
        return
    out[prefix] = value


def _flat_record(row: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    _flatten("", row, out)
    return out


def _write_key_value_sheet(wb: Workbook, name: str, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet(name)
    ws.append(["key", "value"])
    flat = _flat_record(payload)
    for key in sorted(flat):
        value = flat[key]
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        ws.append([key, value])


def _write_records_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(name)
    if not rows:
        ws.append(["empty"])
        return
    flat_rows = [_flat_record(row) for row in rows if isinstance(row, dict)]
    fieldnames = sorted({key for row in flat_rows for key in row.keys()})
    ws.append(fieldnames)
    for row in flat_rows:
        ws.append([
            json.dumps(row.get(field), ensure_ascii=False) if isinstance(row.get(field), (dict, list)) else row.get(field)
            for field in fieldnames
        ])


def _portfolio_rows(portfolio: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for hazard in HAZARDS:
        haz = portfolio.get(hazard) if isinstance(portfolio.get(hazard), dict) else {}
        rows.append({"hazard": hazard, **haz})
    delta = portfolio.get("delta") if isinstance(portfolio.get("delta"), dict) else None
    if delta:
        rows.append({"hazard": "delta", **delta})
    return rows


def _pml_rows(portfolio: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for hazard in HAZARDS:
        haz = portfolio.get(hazard) if isinstance(portfolio.get(hazard), dict) else {}
        rows.append(
            {
                "hazard": hazard,
                "eai_eur": haz.get("eai_eur") or haz.get("aai_agg_eur"),
                "pml_50_eur": haz.get("pml_50_eur"),
                "pml_100_eur": haz.get("pml_100_eur"),
                "pml_1000_eur": haz.get("pml_1000_eur"),
            }
        )
    return rows


def _network_rows(result: dict[str, Any]) -> list[dict[str, Any]]:
    payload = (
        result.get("pml_network_graph_inputs", {}).get("network_state_service_distribution_by_scenario")
        if isinstance(result.get("pml_network_graph_inputs"), dict)
        else None
    )
    rows: list[dict[str, Any]] = []
    if isinstance(payload, dict):
        for scenario, hazards in payload.items():
            if not isinstance(hazards, dict):
                continue
            for hazard, groups in hazards.items():
                if not isinstance(groups, dict):
                    continue
                for group, values in groups.items():
                    if isinstance(values, dict):
                        rows.append({"scenario": scenario, "hazard": hazard, "service_group": group, **values})
    portfolio = result.get("portfolio_results") if isinstance(result.get("portfolio_results"), dict) else {}
    projected = portfolio.get("network_states_projected") if isinstance(portfolio.get("network_states_projected"), dict) else None
    if not rows and projected:
        for hazard, groups in projected.items():
            if isinstance(groups, dict):
                for group, values in groups.items():
                    if isinstance(values, dict):
                        rows.append({"scenario": "rp100", "hazard": hazard, "service_group": group, **values})
    return rows


def _input_feature_rows(result: dict[str, Any]) -> list[dict[str, Any]]:
    feature_collection = result.get("input_features_geojson")
    features = feature_collection.get("features") if isinstance(feature_collection, dict) else None
    rows: list[dict[str, Any]] = []
    for feature in features if isinstance(features, list) else []:
        if not isinstance(feature, dict):
            continue
        properties = feature.get("properties") if isinstance(feature.get("properties"), dict) else {}
        geometry = feature.get("geometry") if isinstance(feature.get("geometry"), dict) else None
        rows.append(
            {
                **properties,
                "geometry_geojson": json.dumps(geometry, ensure_ascii=False) if geometry else None,
            }
        )
    return rows


def build_xlsx_bytes(result: dict[str, Any]) -> bytes:
    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)
    _write_key_value_sheet(wb, "meta", result.get("meta") if isinstance(result.get("meta"), dict) else {})
    _write_key_value_sheet(
        wb,
        "exposure_summary",
        result.get("exposure_summary") if isinstance(result.get("exposure_summary"), dict) else {},
    )
    _write_records_sheet(wb, "asset_results", result.get("asset_results") if isinstance(result.get("asset_results"), list) else [])
    _write_records_sheet(wb, "input_features", _input_feature_rows(result))
    _write_records_sheet(
        wb,
        "territory_results",
        result.get("territory_results") if isinstance(result.get("territory_results"), list) else [],
    )
    portfolio = result.get("portfolio_results") if isinstance(result.get("portfolio_results"), dict) else {}
    _write_records_sheet(wb, "portfolio_results", _portfolio_rows(portfolio))
    _write_records_sheet(wb, "pml_scenarios", _pml_rows(portfolio))
    _write_records_sheet(wb, "network_states", _network_rows(result))
    social = portfolio.get("social_impact_summary") if isinstance(portfolio.get("social_impact_summary"), dict) else {}
    _write_key_value_sheet(wb, "social_impacts", social)
    event_summary = portfolio.get("event_summary") if isinstance(portfolio.get("event_summary"), dict) else {}
    event_rows: list[dict[str, Any]] = []
    for hazard, events in event_summary.items():
        if isinstance(events, list):
            for event in events:
                if isinstance(event, dict):
                    event_rows.append({"hazard": hazard, **event})
    _write_records_sheet(wb, "event_summary", event_rows)
    downloads = result.get("artifacts", {}).get("downloads") if isinstance(result.get("artifacts"), dict) else []
    _write_records_sheet(wb, "graph_manifest", downloads if isinstance(downloads, list) else [])
    notes = result.get("notes") if isinstance(result.get("notes"), list) else []
    _write_records_sheet(wb, "warnings", [{"warning": str(item)} for item in notes])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _load_matplotlib():
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    return plt, PdfPages


def _chart_values(result: dict[str, Any], scenario: str) -> tuple[list[str], list[float]]:
    portfolio = result.get("portfolio_results") if isinstance(result.get("portfolio_results"), dict) else {}
    labels = ["STORM", "STORM_CMCC"]
    key_by_scenario = {
        "annual": "eai_eur",
        "rp50": "pml_50_eur",
        "rp100": "pml_100_eur",
        "rp1000": "pml_1000_eur",
    }
    key = key_by_scenario.get(scenario, "eai_eur")
    values = [
        _num((portfolio.get("storm") or {}).get(key)),
        _num((portfolio.get("storm_cmcc") or {}).get(key)),
    ]
    return labels, values


def _render_bar_chart(path: Path, title: str, labels: list[str], values: list[float]) -> None:
    plt, _ = _load_matplotlib()
    fig, ax = plt.subplots(figsize=(9, 5))
    colors = ["#2563eb", "#f59e0b"]
    ax.bar(labels, [v / 1_000_000.0 for v in values], color=colors[: len(labels)])
    ax.set_ylabel("MEUR")
    ax.set_title(title)
    ax.grid(axis="y", alpha=0.25)
    for idx, value in enumerate(values):
        ax.text(idx, value / 1_000_000.0, f"{value / 1_000_000.0:.2f}", ha="center", va="bottom", fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _render_state_chart(path: Path, result: dict[str, Any]) -> bool:
    rows = _network_rows(result)
    if not rows:
        return False
    totals = {"S0": 0.0, "S1": 0.0, "S2": 0.0, "S3": 0.0}
    for row in rows:
        for state in totals:
            totals[state] += _num(row.get(state))
    plt, _ = _load_matplotlib()
    fig, ax = plt.subplots(figsize=(9, 5))
    labels = list(totals.keys())
    values = [totals[key] for key in labels]
    ax.bar(labels, values, color=["#6AB96F", "#f2b66f", "#d94832", "#111111"])
    ax.set_title("Etats reseau projetes")
    ax.set_ylabel("Unites")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return True


def _load_basin_overlays() -> dict[str, Any]:
    if not BASIN_OVERLAYS_PATH.exists():
        return {}
    try:
        payload = json.loads(BASIN_OVERLAYS_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}
    basins = payload.get("basins") if isinstance(payload, dict) else None
    return basins if isinstance(basins, dict) else {}


def _feature_collection(result: dict[str, Any]) -> list[dict[str, Any]]:
    collection = result.get("input_features_geojson")
    features = collection.get("features") if isinstance(collection, dict) else None
    return [feature for feature in features if isinstance(feature, dict)] if isinstance(features, list) else []


def _asset_results_by_id(result: dict[str, Any]) -> dict[str, dict[str, Any]]:
    rows = result.get("asset_results") if isinstance(result.get("asset_results"), list) else []
    return {
        str(row.get("asset_id") or ""): row
        for row in rows
        if isinstance(row, dict) and str(row.get("asset_id") or "")
    }


def _iter_positions(value: Any):
    if not isinstance(value, (list, tuple)):
        return
    if len(value) >= 2 and all(isinstance(item, (int, float)) and math.isfinite(float(item)) for item in value[:2]):
        yield float(value[0]), float(value[1])
        return
    for child in value:
        yield from _iter_positions(child)


def _geometry_positions(geometry: Any) -> list[tuple[float, float]]:
    if not isinstance(geometry, dict):
        return []
    if str(geometry.get("type") or "") == "GeometryCollection":
        out: list[tuple[float, float]] = []
        for child in geometry.get("geometries") or []:
            out.extend(_geometry_positions(child))
        return out
    return list(_iter_positions(geometry.get("coordinates")))


def _feature_basin(feature: dict[str, Any], basins: dict[str, Any]) -> str | None:
    properties = feature.get("properties") if isinstance(feature.get("properties"), dict) else {}
    explicit = str(properties.get("quick_basin") or "").strip().lower()
    if explicit in basins:
        return explicit
    positions = _geometry_positions(feature.get("geometry"))
    if not positions:
        return None
    lon = sum(item[0] for item in positions) / len(positions)
    lat = sum(item[1] for item in positions) / len(positions)
    for basin, payload in basins.items():
        bounds = payload.get("bounds") if isinstance(payload, dict) else None
        if not isinstance(bounds, dict):
            continue
        if (
            _num(bounds.get("west")) <= lon <= _num(bounds.get("east"))
            and _num(bounds.get("south")) <= lat <= _num(bounds.get("north"))
        ):
            return str(basin)
    return None


def _features_by_basin(features: list[dict[str, Any]], basins: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for feature in features:
        basin = _feature_basin(feature, basins)
        if basin:
            grouped.setdefault(basin, []).append(feature)
    return grouped


def _local_extent(features: list[dict[str, Any]], basin_payload: dict[str, Any] | None = None) -> tuple[float, float, float, float]:
    positions = [position for feature in features for position in _geometry_positions(feature.get("geometry"))]
    bounds = basin_payload.get("bounds") if isinstance(basin_payload, dict) else None
    if not positions and isinstance(bounds, dict):
        return (
            _num(bounds.get("west")),
            _num(bounds.get("east")),
            _num(bounds.get("south")),
            _num(bounds.get("north")),
        )
    if not positions:
        return (-180.0, 180.0, -60.0, 60.0)
    xs = [item[0] for item in positions]
    ys = [item[1] for item in positions]
    center_x = (min(xs) + max(xs)) / 2.0
    center_y = (min(ys) + max(ys)) / 2.0
    span_x = max(max(xs) - min(xs), 3.0)
    span_y = max(max(ys) - min(ys), 3.0)
    span = max(span_x, span_y)
    west, east = center_x - span * 0.65, center_x + span * 0.65
    south, north = center_y - span * 0.65, center_y + span * 0.65
    if isinstance(bounds, dict):
        west = max(west, _num(bounds.get("west")))
        east = min(east, _num(bounds.get("east")))
        south = max(south, _num(bounds.get("south")))
        north = min(north, _num(bounds.get("north")))
    return west, east, south, north


def _plot_geometry(ax: Any, geometry: Any, *, color: str, label: str | None = None, alpha: float = 1.0, zorder: int = 5) -> None:
    if not isinstance(geometry, dict):
        return
    geometry_type = str(geometry.get("type") or "")
    coordinates = geometry.get("coordinates")
    if geometry_type == "Point":
        positions = _geometry_positions(geometry)
        if positions:
            ax.scatter([positions[0][0]], [positions[0][1]], s=46, color=color, edgecolor="white", linewidth=0.8, alpha=alpha, zorder=zorder, label=label)
        return
    if geometry_type == "MultiPoint":
        positions = _geometry_positions(geometry)
        if positions:
            ax.scatter([item[0] for item in positions], [item[1] for item in positions], s=38, color=color, edgecolor="white", linewidth=0.7, alpha=alpha, zorder=zorder, label=label)
        return
    if geometry_type == "LineString":
        positions = list(_iter_positions(coordinates))
        if positions:
            ax.plot([item[0] for item in positions], [item[1] for item in positions], color=color, linewidth=3.0, alpha=alpha, solid_capstyle="round", zorder=zorder, label=label)
        return
    if geometry_type == "Polygon":
        rings = coordinates if isinstance(coordinates, list) else []
        for index, ring in enumerate(rings):
            positions = list(_iter_positions(ring))
            if not positions:
                continue
            xs = [item[0] for item in positions]
            ys = [item[1] for item in positions]
            if index == 0:
                ax.fill(xs, ys, color=color, alpha=min(alpha, 0.28), zorder=zorder)
            ax.plot(xs, ys, color=color, linewidth=2.2, alpha=alpha, zorder=zorder + 1, label=label if index == 0 else None)
        return
    if geometry_type.startswith("Multi"):
        singular_type = geometry_type.removeprefix("Multi")
        for index, child in enumerate(coordinates if isinstance(coordinates, list) else []):
            _plot_geometry(
                ax,
                {"type": singular_type, "coordinates": child},
                color=color,
                label=label if index == 0 else None,
                alpha=alpha,
                zorder=zorder,
            )
        return
    if geometry_type == "GeometryCollection":
        for index, child in enumerate(geometry.get("geometries") or []):
            _plot_geometry(ax, child, color=color, label=label if index == 0 else None, alpha=alpha, zorder=zorder)


def _style_geo_axis(ax: Any, extent: tuple[float, float, float, float]) -> None:
    west, east, south, north = extent
    ax.set_xlim(west, east)
    ax.set_ylim(south, north)
    ax.set_facecolor("#dce3e6")
    ax.set_xlabel("Longitude")
    ax.set_ylabel("Latitude")
    ax.grid(color="white", linewidth=0.7, alpha=0.75)
    ax.set_aspect("equal", adjustable="box")


def _add_light_basemap(ax: Any) -> None:
    """Add the same neutral Esri context used by the complete graph generator."""
    try:
        import contextily as cx

        cx.add_basemap(
            ax,
            crs="EPSG:4326",
            source=cx.providers.Esri.WorldGrayCanvas,
            attribution=False,
            zoom="auto",
            reset_extent=False,
        )
    except Exception:
        # Report generation must remain available when the tile service is offline.
        return


def _render_hazard_map(path: Path, result: dict[str, Any]) -> bool:
    features = _feature_collection(result)
    basins = _load_basin_overlays()
    grouped = _features_by_basin(features, basins)
    if not grouped:
        return False
    plt, _ = _load_matplotlib()
    from matplotlib.cm import ScalarMappable
    from matplotlib.colors import Normalize

    basin_items = sorted(grouped.items())
    fig, axes = plt.subplots(len(basin_items), len(HAZARDS), figsize=(14, 5.8 * len(basin_items)), squeeze=False)
    for row_index, (basin, basin_features) in enumerate(basin_items):
        basin_payload = basins.get(basin) if isinstance(basins.get(basin), dict) else {}
        bounds = basin_payload.get("bounds") if isinstance(basin_payload.get("bounds"), dict) else {}
        extent = _local_extent(basin_features, basin_payload)
        raster_extent = [
            _num(bounds.get("west")),
            _num(bounds.get("east")),
            _num(bounds.get("south")),
            _num(bounds.get("north")),
        ]
        metric = basin_payload.get("metrics", {}).get("rp100", {}) if isinstance(basin_payload.get("metrics"), dict) else {}
        metric_min = _num(metric.get("min_mps"))
        metric_max = _num(metric.get("max_mps"))
        for col_index, hazard in enumerate(HAZARDS):
            ax = axes[row_index][col_index]
            _style_geo_axis(ax, extent)
            _add_light_basemap(ax)
            overlay_rel = basin_payload.get("overlays", {}).get(hazard, {}).get("rp100") if isinstance(basin_payload.get("overlays"), dict) else None
            overlay_path = BASIN_OVERLAYS_PATH.parent / str(overlay_rel or "")
            if overlay_rel and overlay_path.exists():
                ax.imshow(plt.imread(str(overlay_path)), extent=raster_extent, origin="upper", interpolation="nearest", zorder=1)
            for feature in basin_features:
                _plot_geometry(ax, feature.get("geometry"), color="#c2185b", alpha=1.0, zorder=5)
            climate = "Climat actuel (STORM)" if hazard == "storm" else "Climat 2050 (STORM_CMCC)"
            ax.set_title(f"{str(basin_payload.get('label') or basin).title()} - {climate}\nVent cyclonique RP100")
            if metric_max > metric_min:
                colorbar = fig.colorbar(
                    ScalarMappable(norm=Normalize(vmin=metric_min, vmax=metric_max), cmap="turbo"),
                    ax=ax,
                    orientation="horizontal",
                    fraction=0.046,
                    pad=0.10,
                )
                colorbar.set_label("Vitesse du vent (m/s)")
    fig.suptitle("Carte des aleas sur l'emprise utilisateur", fontsize=16)
    fig.text(0.5, 0.01, "Rasters de bassin SIB Work; geometries utilisateur en rose. Fond Esri WorldGrayCanvas si disponible.", ha="center", fontsize=9)
    fig.tight_layout(rect=(0, 0.03, 1, 0.96))
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return True


def _state_from_ratio(loss_ratio: float) -> str:
    for state, threshold in STATE_THRESHOLDS:
        if loss_ratio >= threshold:
            return state
    return "S0"


def _asset_state(result: dict[str, Any], row: dict[str, Any], scenario: str, hazard: str) -> tuple[str, str]:
    direct_state = str(row.get(f"service_state_{scenario}_{hazard}") or "").upper()
    if direct_state in STATE_COLORS:
        return direct_state, str(row.get(f"service_state_source_{scenario}_{hazard}") or "quick_impact_state")
    exposure = max(_num(row.get("exposure_eur")), 1.0)
    scenario_losses = row.get("scenario_losses_eur") if isinstance(row.get("scenario_losses_eur"), dict) else {}
    loss = _num((scenario_losses.get(hazard) or {}).get(scenario)) if isinstance(scenario_losses.get(hazard), dict) else 0.0
    if loss > 0:
        return _state_from_ratio(loss / exposure), "asset_scenario_loss"
    portfolio = result.get("portfolio_results") if isinstance(result.get("portfolio_results"), dict) else {}
    hazard_portfolio = portfolio.get(hazard) if isinstance(portfolio.get(hazard), dict) else {}
    suffix = "storm" if hazard == "storm" else "cmcc"
    asset_eai = _num(row.get(f"eai_{suffix}_eur"))
    portfolio_eai = _num(hazard_portfolio.get("eai_eur") or hazard_portfolio.get("aai_agg_eur"))
    rp_key = "pml_100_eur" if scenario == "rp100" else "pml_1000_eur"
    multiplier = _num(hazard_portfolio.get(rp_key)) / portfolio_eai if portfolio_eai > 0 else 1.0
    return _state_from_ratio((asset_eai * multiplier) / exposure), "portfolio_pml_scaled_asset_eai"


def _render_network_state_map(path: Path, result: dict[str, Any], scenario: str) -> bool:
    features = _feature_collection(result)
    basins = _load_basin_overlays()
    grouped = _features_by_basin(features, basins)
    if not grouped:
        return False
    rows_by_id = _asset_results_by_id(result)
    plt, _ = _load_matplotlib()
    from matplotlib.lines import Line2D

    basin_items = sorted(grouped.items())
    fig, axes = plt.subplots(len(basin_items), len(HAZARDS), figsize=(14, 5.8 * len(basin_items)), squeeze=False)
    source_labels: set[str] = set()
    for row_index, (basin, basin_features) in enumerate(basin_items):
        basin_payload = basins.get(basin) if isinstance(basins.get(basin), dict) else {}
        extent = _local_extent(basin_features, basin_payload)
        for col_index, hazard in enumerate(HAZARDS):
            ax = axes[row_index][col_index]
            _style_geo_axis(ax, extent)
            _add_light_basemap(ax)
            for feature in basin_features:
                properties = feature.get("properties") if isinstance(feature.get("properties"), dict) else {}
                asset_id = str(properties.get("asset_id") or "")
                row = rows_by_id.get(asset_id, properties)
                state, source = _asset_state(result, row, scenario, hazard)
                source_labels.add(source)
                _plot_geometry(ax, feature.get("geometry"), color=STATE_COLORS[state], alpha=1.0, zorder=5)
            climate = "Climat actuel (STORM)" if hazard == "storm" else "Climat 2050 (STORM_CMCC)"
            period = "100" if scenario == "rp100" else "1 000"
            ax.set_title(f"{str(basin_payload.get('label') or basin).title()} - {climate}\nEtat des infrastructures - retour {period} ans")
            ax.legend(
                handles=[Line2D([0], [0], color=color, lw=4, marker="o", markersize=5, label=state) for state, color in STATE_COLORS.items()],
                title="Etat de service",
                loc="upper right",
                framealpha=0.94,
            )
    period = "100" if scenario == "rp100" else "1 000"
    fig.suptitle(f"Carte de defaillance des infrastructures dessinees - RP{period.replace(' ', '')}", fontsize=16)
    if "portfolio_pml_scaled_asset_eai" in source_labels:
        note = "Etat estime par actif a partir de l'EAI et du PML portefeuille pour les sorties completes sans etat spatial individuel."
    else:
        note = "Etats S0-S3 issus des resultats par actif du moteur de calcul."
    fig.text(0.5, 0.01, note, ha="center", fontsize=9)
    fig.tight_layout(rect=(0, 0.03, 1, 0.96))
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return True


def render_graph_pack(result: dict[str, Any], output_dir: Path) -> dict[str, Any]:
    graphs_dir = output_dir
    graphs_dir.mkdir(parents=True, exist_ok=True)
    generated: list[dict[str, Any]] = []
    hazard_map_path = graphs_dir / "user-hazard-map-rp100.png"
    if _render_hazard_map(hazard_map_path, result):
        generated.append(
            {
                "path": hazard_map_path.relative_to(output_dir).as_posix(),
                "technical_type": "map",
                "family": "hazard",
                "graph_type": "hazard_map_rp100",
                "title": "Carte des aleas vent cyclonique - RP100",
            }
        )
    for scenario, label in SCENARIOS:
        chart_path = graphs_dir / f"user-{scenario}.png"
        labels, values = _chart_values(result, scenario)
        _render_bar_chart(chart_path, f"{label} - STORM vs STORM_CMCC", labels, values)
        generated.append(
            {
                "path": chart_path.relative_to(output_dir).as_posix(),
                "technical_type": "chart",
                "family": "impact",
                "graph_type": scenario,
                "title": f"{label} - STORM vs STORM_CMCC",
            }
        )
    state_path = graphs_dir / "user-network-states.png"
    if _render_state_chart(state_path, result):
        generated.append(
            {
                "path": state_path.relative_to(output_dir).as_posix(),
                "technical_type": "chart",
                "family": "network",
                "graph_type": "network_states",
                "title": "Etats reseau projetes",
            }
        )
    for scenario, period in (("rp100", "100"), ("rp1000", "1000")):
        state_map_path = graphs_dir / f"user-network-state-map-rp{period}.png"
        if _render_network_state_map(state_map_path, result, scenario):
            generated.append(
                {
                    "path": state_map_path.relative_to(output_dir).as_posix(),
                    "technical_type": "map",
                    "family": "network",
                    "graph_type": f"network_state_map_{scenario}",
                    "title": f"Carte de defaillance des infrastructures - RP{period}",
                }
            )
    manifest = {
        "schema_version": "sib_user_graph_manifest_v1",
        "generated_at": datetime.now(UTC).replace(microsecond=0).isoformat(),
        "source": "user_result_payload",
        "generated_outputs": generated,
    }
    (output_dir / "graphs-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def _write_pdf(path: Path, result: dict[str, Any], manifest: dict[str, Any], output_dir: Path) -> None:
    plt, PdfPages = _load_matplotlib()
    with PdfPages(path) as pdf:
        fig = plt.figure(figsize=(8.27, 11.69))
        ax = fig.add_subplot(111)
        ax.axis("off")
        meta = result.get("meta") if isinstance(result.get("meta"), dict) else {}
        summary = result.get("exposure_summary") if isinstance(result.get("exposure_summary"), dict) else {}
        lines = [
            "Rapport SIB - donnees utilisateurs",
            "",
            f"Run: {meta.get('run_label') or meta.get('job_id') or ''}",
            f"Moteur: {meta.get('engine') or ''}",
            f"Reference: {meta.get('reference_engine') or meta.get('impact_function') or ''}",
            f"Date: {meta.get('updated_at') or ''}",
            "",
            f"Actifs: {summary.get('asset_count_original') or summary.get('asset_count_points') or 0}",
            f"Exposition totale EUR: {_num(summary.get('total_exposure_eur')):,.0f}",
        ]
        notes = result.get("notes") if isinstance(result.get("notes"), list) else []
        if notes:
            lines.extend(["", "Limites et avertissements:"])
            lines.extend(f"- {str(note)[:180]}" for note in notes[:8])
        ax.text(0.08, 0.92, "\n".join(lines), ha="left", va="top", fontsize=11, wrap=True)
        pdf.savefig(fig)
        plt.close(fig)

        for item in manifest.get("generated_outputs") or []:
            rel_path = item.get("path")
            if not rel_path:
                continue
            image_path = output_dir / str(rel_path)
            if not image_path.exists():
                continue
            fig = plt.figure(figsize=(11.69, 8.27))
            ax = fig.add_subplot(111)
            ax.axis("off")
            img = plt.imread(str(image_path))
            ax.imshow(img)
            ax.set_title(str(item.get("title") or ""), fontsize=12)
            pdf.savefig(fig)
            plt.close(fig)


def artifact_kind_for_name(name: str) -> str:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".xlsx":
        return "xlsx"
    if suffix == ".csv":
        return "csv"
    if suffix == ".geojson":
        return "geojson"
    if suffix == ".json":
        return "json"
    if suffix == ".html":
        return "html"
    if suffix == ".png":
        return "png"
    return "file"


def mime_type_for_name(name: str) -> str:
    if name.endswith(".xlsx"):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    guessed = mimetypes.guess_type(name)[0]
    return guessed or "application/octet-stream"


def ensure_output_artifacts(job_id: str, result: dict[str, Any], store: JobStore) -> list[Path]:
    output_dir = store.artifacts_dir(job_id)
    manifest = render_graph_pack(result, output_dir)
    generated_paths = [
        output_dir / str(item.get("path") or "")
        for item in manifest.get("generated_outputs") or []
        if isinstance(item, dict) and item.get("path")
    ]
    artifacts = result.setdefault("artifacts", {})
    artifacts["visuals"] = [
        {
            "name": path.name,
            "url": f"/api/v1/runs/{job_id}/artifacts/{path.name}",
            "graph_type": item.get("graph_type"),
            "technical_type": item.get("technical_type"),
            "title": item.get("title"),
        }
        for item, path in zip(manifest.get("generated_outputs") or [], generated_paths)
        if path.exists()
    ]
    artifacts["plots_png"] = [path.name for path in generated_paths if path.exists()]
    artifacts["downloads"] = [
        {"name": "user-results-report.pdf", "url": f"/api/v1/runs/{job_id}/artifacts/user-results-report.pdf"},
        {"name": "user-results-matrix.xlsx", "url": f"/api/v1/runs/{job_id}/artifacts/user-results-matrix.xlsx"},
    ]
    xlsx_path = store.save_artifact_bytes(job_id, "user-results-matrix.xlsx", build_xlsx_bytes(result))
    pdf_path = output_dir / "user-results-report.pdf"
    _write_pdf(pdf_path, result, manifest, output_dir)
    graph_manifest_path = output_dir / "graphs-manifest.json"
    store.save_result(job_id, result)
    return [xlsx_path, pdf_path, graph_manifest_path, *generated_paths]
