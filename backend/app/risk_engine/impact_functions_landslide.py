from __future__ import annotations

from pathlib import Path
from typing import Any

try:
    import numpy as np  # type: ignore
except Exception:  # pragma: no cover - optional during import
    np = None  # type: ignore

from .impact_functions import ASSET_TYPE_TO_CURVE_CODE
from .impact_functions_multi_hazard import FLOOD_ASSET_TYPE_TO_CURVE_CODE


LANDSLIDE_IMPF_ID_HYPOTHESIS = 6100
LANDSLIDE_IMPF_ID_D2_PROXY = 6101
LANDSLIDE_HAZ_TYPE = "LS"
LANDSLIDE_INTENSITY_UNIT = "class"
LANDSLIDE_INTENSITIES = [0.0, 1.0, 2.0, 3.0, 4.0, 5.0]
LANDSLIDE_HYPOTHESIS_CURVE = [0.0, 0.0, 0.25, 0.50, 0.75, 1.0]
LANDSLIDE_PROXY_FALLBACK_CURVE = [0.0, 0.0, 0.15, 0.30, 0.76, 1.0]
LANDSLIDE_D2_WORKBOOK = Path("/home/ubuntu/uploads/Vulnerability/Table_D2_Hazard_Fragility_and_Vulnerability_Curves_V1.1.0.xlsx")

_CURVE_CACHE: dict[str, dict[str, Any]] | None = None


def _all_sib_asset_types() -> list[str]:
    asset_types = set(ASSET_TYPE_TO_CURVE_CODE.keys()) | set(FLOOD_ASSET_TYPE_TO_CURVE_CODE.keys())
    return sorted(str(asset) for asset in asset_types)


def _build_curve_entry(
    *,
    impf_id: int,
    code: str,
    name: str,
    source: str,
    geography: str,
    mdd: list[float],
    modeled_infrastructure_type: str,
    modeled_infrastructure_characteristics: str,
    proxy_source: str | None = None,
) -> dict[str, Any]:
    return {
        "impf_id": int(impf_id),
        "code": str(code),
        "name": str(name),
        "source": str(source),
        "geography": str(geography),
        "haz_type": LANDSLIDE_HAZ_TYPE,
        "intensity_unit": LANDSLIDE_INTENSITY_UNIT,
        "intensity": [float(v) for v in LANDSLIDE_INTENSITIES],
        "mdd": [float(v) for v in mdd],
        "paa": [1.0 for _ in LANDSLIDE_INTENSITIES],
        "modeled_infrastructure_type": str(modeled_infrastructure_type),
        "modeled_infrastructure_characteristics": str(modeled_infrastructure_characteristics),
        "uncertainty_lower": None,
        "uncertainty_upper": None,
        "proxy_source": proxy_source,
    }


def _extract_row_average_from_categorical_sheet(ws: Any, target_label: str) -> float | None:
    for row in ws.iter_rows(values_only=True):
        values = list(row)
        if target_label not in {str(v).strip() for v in values if v is not None}:
            continue
        try:
            idx = next(i for i, value in enumerate(values) if str(value).strip() == target_label)
        except StopIteration:
            continue
        numeric_values = [float(value) for value in values[idx + 1 :] if isinstance(value, (int, float))]
        if not numeric_values:
            return None
        return float(sum(numeric_values) / len(numeric_values))
    return None


def _build_d2_proxy_curve_from_workbook(workbook_path: Path | None = None) -> list[float]:
    fallback = list(LANDSLIDE_PROXY_FALLBACK_CURVE)
    path = Path(workbook_path or LANDSLIDE_D2_WORKBOOK)
    if not path.exists():
        return fallback

    try:
        from openpyxl import load_workbook
    except Exception:
        return fallback

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
        if "L_Categorical" not in wb.sheetnames:
            return fallback
        ws = wb["L_Categorical"]
        class_to_label = {
            2.0: "M-I",
            3.0: "M-II",
            4.0: "M-III",
        }
        values = [0.0, 0.0]
        for class_value in (2.0, 3.0, 4.0):
            label = class_to_label[class_value]
            avg = _extract_row_average_from_categorical_sheet(ws, label)
            if avg is None:
                return fallback
            values.append(max(0.0, min(1.0, float(avg))))
        values.append(1.0)
    except Exception:
        return fallback

    if np is not None:
        arr = np.asarray(values, dtype=float)
        arr = np.maximum.accumulate(arr)
        arr[0] = 0.0
        arr[1] = 0.0
        arr[-1] = 1.0
        return [float(v) for v in arr.tolist()]

    out: list[float] = []
    running = 0.0
    for idx, value in enumerate(values):
        running = max(running, float(value))
        if idx in {0, 1}:
            running = 0.0
        out.append(running)
    out[-1] = 1.0
    return out


def _build_curve_catalog(workbook_path: Path | None = None) -> dict[int, dict[str, Any]]:
    asset_types = _all_sib_asset_types()
    proxy_curve = _build_d2_proxy_curve_from_workbook(workbook_path)
    default_curve = _build_curve_entry(
        impf_id=LANDSLIDE_IMPF_ID_HYPOTHESIS,
        code="LS_HYPOTHESIS_5STEP",
        name="Hypothèse 5 paliers",
        source="Hypothèse SIB",
        geography="Guadeloupe / Martinique",
        mdd=list(LANDSLIDE_HYPOTHESIS_CURVE),
        modeled_infrastructure_type="Portefeuille SIB",
        modeled_infrastructure_characteristics="Generic probabilistic landslide curve",
    )
    proxy_entry = _build_curve_entry(
        impf_id=LANDSLIDE_IMPF_ID_D2_PROXY,
        code="LS_D2_PROXY",
        name="Proxy D2",
        source="D2 workbook L_Categorical",
        geography="Proxy derived from the D2 landslide sheet",
        mdd=proxy_curve,
        modeled_infrastructure_type="Portefeuille SIB",
        modeled_infrastructure_characteristics="Proxy derived from the D2 landslide categorical sheet",
        proxy_source="L_Categorical",
    )
    default_curve["sib_asset_types"] = asset_types
    proxy_entry["sib_asset_types"] = asset_types

    return {
        LANDSLIDE_IMPF_ID_HYPOTHESIS: default_curve,
        LANDSLIDE_IMPF_ID_D2_PROXY: proxy_entry,
    }


def get_landslide_curve_catalog(workbook_path: Path | None = None) -> dict[int, dict[str, Any]]:
    global _CURVE_CACHE
    if _CURVE_CACHE is None:
        _CURVE_CACHE = _build_curve_catalog(workbook_path)
    return _CURVE_CACHE


def get_landslide_vulnerability_payload(*, d2_curve_file: Path | None = None) -> dict[str, Any]:
    catalog = get_landslide_curve_catalog(d2_curve_file)
    curves = [dict(catalog[key]) for key in sorted(catalog.keys())]

    explicit_mapping: dict[str, dict[str, Any]] = {}
    for asset_type in _all_sib_asset_types():
        explicit_mapping[asset_type] = {
            "code": "LS_HYPOTHESIS_5STEP",
            "impf_id": LANDSLIDE_IMPF_ID_HYPOTHESIS,
        }

    return {
        "profile": "sib_landslide_multicurve_v1",
        "haz_type": LANDSLIDE_HAZ_TYPE,
        "hazard_component": "landslide",
        "intensity_unit": LANDSLIDE_INTENSITY_UNIT,
        "default_curve": {
            "code": "LS_HYPOTHESIS_5STEP",
            "impf_id": LANDSLIDE_IMPF_ID_HYPOTHESIS,
        },
        "explicit_asset_type_mapping": explicit_mapping,
        "curves": curves,
    }


def _build_climada_impact_func(curve: dict[str, Any]) -> Any | None:
    if np is None:
        return None
    try:
        from climada.entity.impact_funcs.base import ImpactFunc  # type: ignore
    except Exception:
        return None

    impf = ImpactFunc()
    impf.id = int(curve["impf_id"])
    impf.haz_type = str(curve.get("haz_type") or LANDSLIDE_HAZ_TYPE)
    impf.name = str(curve.get("name") or f"LS Impact Function {impf.id}")
    impf.intensity = np.asarray(list(curve.get("intensity") or []), dtype=float)
    impf.mdd = np.clip(np.asarray(list(curve.get("mdd") or []), dtype=float), 0.0, 1.0)
    impf.paa = np.asarray(list(curve.get("paa") or []), dtype=float)
    impf.intensity_unit = str(curve.get("intensity_unit") or LANDSLIDE_INTENSITY_UNIT)
    try:
        impf.check()
    except Exception:
        return None
    return impf


def try_build_climada_landslide_impact_funcs() -> list[Any] | None:
    catalog = get_landslide_curve_catalog()
    default_curve = catalog.get(LANDSLIDE_IMPF_ID_HYPOTHESIS)
    if not default_curve:
        return None
    impf = _build_climada_impact_func(default_curve)
    if impf is None:
        return None
    return [impf]
