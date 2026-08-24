#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import sys


REPO_ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = REPO_ROOT / "backend"
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.risk_engine.sensitivity_scenarios import parameter_traceability

DEFAULT_WORKBOOK_PATH = Path("/home/ubuntu/uploads/Sensibility analysis/Analyses sensibilité_2.xlsx")
DEFAULT_SHEET_NAME = "Variable to modify for sensibil"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "config" / "sensitivity"

RAW_CSV_NAME = "workbook-variable-sheet.csv"
RAW_JSON_NAME = "workbook-variable-sheet.json"
PACK_JSON_NAME = "default-scenario-pack.json"
DEFAULT_PACK_EXCLUDED_PARAMETERS = {
    "uplift_by_state",
}


def _clean_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\xa0", " ")
    return text.strip()


def _slugify(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    normalized = normalized.lower()
    normalized = re.sub(r"[^a-z0-9]+", "-", normalized)
    normalized = normalized.strip("-")
    return normalized or "scenario"


def _parse_decimal(value: str) -> float:
    text = _clean_text(value).replace(",", ".")
    match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    if match is None:
        raise ValueError(f"Unable to parse decimal value from '{value}'")
    return float(match.group(0))


def _parse_decimal_triplet(value: str) -> tuple[float, float, float]:
    parts = [_clean_text(part) for part in _clean_text(value).split("/") if _clean_text(part)]
    if len(parts) != 3:
        raise ValueError(f"Expected 3 values, got {len(parts)} in '{value}'")
    return tuple(_parse_decimal(part) for part in parts)  # type: ignore[return-value]


def _parse_percent_triplet(value: str) -> tuple[float, float, float]:
    parts = [_clean_text(part) for part in _clean_text(value).split("/") if _clean_text(part)]
    if len(parts) != 3:
        raise ValueError(f"Expected 3 values, got {len(parts)} in '{value}'")
    return tuple(_parse_decimal(part) / 100.0 for part in parts)  # type: ignore[return-value]


def _parse_percent_quad(value: str) -> tuple[float, float, float, float]:
    parts = [_clean_text(part) for part in _clean_text(value).split("/") if _clean_text(part)]
    if len(parts) != 4:
        raise ValueError(f"Expected 4 values, got {len(parts)} in '{value}'")
    return tuple(_parse_decimal(part) / 100.0 for part in parts)  # type: ignore[return-value]


def _parse_test_values(raw_value: str) -> list[str]:
    cleaned = _clean_text(raw_value)
    if not cleaned:
        return []
    lowered = cleaned.lower()
    if "run de sensibilite a part" in unicodedata.normalize("NFKD", lowered).encode("ascii", "ignore").decode("ascii"):
        return ["manual_profile_required"]
    return [_clean_text(part) for part in cleaned.split("-") if _clean_text(part)]


def _canonical_parameter_key(raw_parameter: str) -> str:
    key = _slugify(raw_parameter)
    aliases = {
        "vulnerability-curves": "vulnerability_curves_profile",
        "runoff-coeff": "runoff_coeff",
        "seuils-d-etat-directs": "direct_state_thresholds",
        "poids-de-sante-health": "health_weights",
        "seuils-health-dependency-state": "dependency_state_thresholds",
        "uplift-by-state": "uplift_by_state",
        "max-dist-inland-km": "max_dist_inland_km",
        "hazard-dynamic-max-tracks": "hazard_dynamic_max_tracks",
        "maille-territoriale-electrique": "territory_grid_deg",
        "sampling-spacing-m-backend": "default_sampling_spacing_m",
        "climada-max-points-per-feature": "climada_max_points_per_feature",
    }
    return aliases.get(key, key.replace("-", "_"))


def _execution_tier_for_parameter(parameter_key: str) -> str:
    if parameter_key in {
        "direct_state_thresholds",
        "health_weights",
        "dependency_state_thresholds",
        "uplift_by_state",
    }:
        return "postprocess_rerun"
    return "full_rerun"


def _build_overrides(parameter_key: str, candidate_value: str) -> tuple[bool, dict[str, Any], list[str]]:
    if parameter_key == "vulnerability_curves_profile":
        return False, {}, [
            "Workbook requests a dedicated vulnerability-curve sensitivity run.",
            "This scenario remains a placeholder until an alternate mapping profile is defined.",
        ]

    settings_overrides: dict[str, Any] = {}
    notes: list[str] = []

    if parameter_key == "runoff_coeff":
        settings_overrides["multi_hazard_rain_base_runoff_coeff"] = _parse_decimal(candidate_value)
    elif parameter_key == "direct_state_thresholds":
        s1, s2, s3 = _parse_percent_triplet(candidate_value)
        settings_overrides.update(
            {
                "interdependency_state_threshold_s0_to_s1": s1,
                "interdependency_state_threshold_s1_to_s2": s2,
                "interdependency_state_threshold_s2_to_s3": s3,
            }
        )
    elif parameter_key == "health_weights":
        w1, w2, w3 = [_parse_decimal(part) for part in _clean_text(candidate_value).split("/") if _clean_text(part)]
        settings_overrides.update(
            {
                "interdependency_health_weight_s1": w1,
                "interdependency_health_weight_s2": w2,
                "interdependency_health_weight_s3": w3,
            }
        )
    elif parameter_key == "dependency_state_thresholds":
        s1, s2, s3 = _parse_percent_triplet(candidate_value)
        settings_overrides.update(
            {
                "interdependency_dependency_state_threshold_s1": s1,
                "interdependency_dependency_state_threshold_s2": s2,
                "interdependency_dependency_state_threshold_s3": s3,
            }
        )
    elif parameter_key == "uplift_by_state":
        s0, s1, s2, s3 = _parse_percent_quad(candidate_value)
        settings_overrides.update(
            {
                "interdependency_uplift_s0": s0,
                "interdependency_uplift_s1": s1,
                "interdependency_uplift_s2": s2,
                "interdependency_uplift_s3": s3,
            }
        )
    elif parameter_key == "max_dist_inland_km":
        settings_overrides["hazard_rain_max_dist_inland_km"] = _parse_decimal(candidate_value)
    elif parameter_key == "hazard_dynamic_max_tracks":
        settings_overrides["hazard_dynamic_max_tracks"] = int(round(_parse_decimal(candidate_value)))
    elif parameter_key == "territory_grid_deg":
        settings_overrides["territory_grid_deg"] = _parse_decimal(candidate_value)
    elif parameter_key == "default_sampling_spacing_m":
        settings_overrides["default_sampling_spacing_m"] = _parse_decimal(candidate_value)
    elif parameter_key == "climada_max_points_per_feature":
        settings_overrides["climada_max_points_per_feature"] = int(round(_parse_decimal(candidate_value)))
    else:
        notes.append(f"No override builder is implemented for parameter '{parameter_key}'.")
        return False, {}, notes

    return True, {"settings": settings_overrides}, notes


def _candidate_matches_current(parameter_key: str, current_value: str, candidate_value: str) -> bool:
    if candidate_value == "manual_profile_required":
        return False
    try:
        if parameter_key == "runoff_coeff":
            return abs(_parse_decimal(current_value) - _parse_decimal(candidate_value)) < 1e-12
        if parameter_key == "direct_state_thresholds":
            return _parse_percent_triplet(current_value) == _parse_percent_triplet(candidate_value)
        if parameter_key == "health_weights":
            return _parse_decimal_triplet(current_value) == _parse_decimal_triplet(candidate_value)
        if parameter_key == "dependency_state_thresholds":
            return _parse_decimal_triplet(current_value) == _parse_decimal_triplet(candidate_value)
        if parameter_key == "uplift_by_state":
            return _parse_percent_quad(current_value) == _parse_percent_quad(candidate_value)
        if parameter_key == "max_dist_inland_km":
            return abs(_parse_decimal(current_value) - _parse_decimal(candidate_value)) < 1e-12
        if parameter_key == "hazard_dynamic_max_tracks":
            return int(round(_parse_decimal(current_value))) == int(round(_parse_decimal(candidate_value)))
        if parameter_key == "territory_grid_deg":
            return abs(_parse_decimal(current_value) - _parse_decimal(candidate_value)) < 1e-12
        if parameter_key == "default_sampling_spacing_m":
            return abs(_parse_decimal(current_value) - _parse_decimal(candidate_value)) < 1e-12
        if parameter_key == "climada_max_points_per_feature":
            return int(round(_parse_decimal(current_value))) == int(round(_parse_decimal(candidate_value)))
    except ValueError:
        return False
    return False


def _read_sheet_rows(workbook_path: Path, sheet_name: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in {workbook_path}; available: {workbook.sheetnames}")

    sheet = workbook[sheet_name]
    header = [_clean_text(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    rows: list[dict[str, Any]] = []
    workbook_total_runs = None

    for row_values in sheet.iter_rows(min_row=2, values_only=True):
        row = {header[idx]: row_values[idx] for idx in range(len(header))}
        if workbook_total_runs is None and row.get("Nombre total de run") is not None:
            workbook_total_runs = int(row.get("Nombre total de run"))
        if not any(value is not None and _clean_text(value) for value in row.values()):
            continue
        rows.append(row)

    return rows, {
        "sheet_name": sheet_name,
        "workbook_total_runs": workbook_total_runs,
        "sheet_max_row": sheet.max_row,
        "sheet_max_column": sheet.max_column,
    }


def build_outputs(workbook_path: Path, sheet_name: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    raw_rows, sheet_meta = _read_sheet_rows(workbook_path, sheet_name)
    variables: list[dict[str, Any]] = []
    scenarios: list[dict[str, Any]] = [
        {
            "scenario_id": "all-default",
            "label": "All default",
            "parameter_key": None,
            "execution_tier": "full_rerun",
            "supported": True,
            "overrides": {"settings": {}},
            "notes": ["Reference scenario with current default parameters."],
        }
    ]

    for raw_row in raw_rows:
        parameter_raw = _clean_text(raw_row.get("parametre"))
        parameter_key = _canonical_parameter_key(parameter_raw)
        current_value = _clean_text(raw_row.get("valeur actuelle"))
        test_values = _parse_test_values(_clean_text(raw_row.get("Set de valeurs à tester")))
        default_pack_test_values = [
            candidate_value
            for candidate_value in test_values
            if not _candidate_matches_current(parameter_key, current_value, candidate_value)
        ]
        if parameter_key in DEFAULT_PACK_EXCLUDED_PARAMETERS:
            default_pack_test_values = []
        run_count = raw_row.get("Nombre de run")
        variable_entry = {
            "category": _clean_text(raw_row.get("Categorie")),
            "parameter_key": parameter_key,
            "parameter_label": parameter_raw,
            "current_value": current_value,
            "role": _clean_text(raw_row.get("role")),
            "priority": _clean_text(raw_row.get("priorite")),
            "robustness": _clean_text(raw_row.get("robustesse")),
            "source": _clean_text(raw_row.get("source")),
            "recommended_action": _clean_text(raw_row.get("action recommandee")),
            "test_values_raw": _clean_text(raw_row.get("Set de valeurs à tester")),
            "test_values": test_values,
            "default_pack_test_values": default_pack_test_values,
            "run_count": None if run_count is None else int(run_count),
            "default_pack_run_count": len(default_pack_test_values),
            "execution_tier": _execution_tier_for_parameter(parameter_key),
            "traceability": parameter_traceability(parameter_key),
        }
        variables.append(variable_entry)

        for candidate_value in default_pack_test_values:
            supported, overrides, notes = _build_overrides(parameter_key, candidate_value)
            scenarios.append(
                {
                    "scenario_id": f"{parameter_key}-{_slugify(candidate_value)}",
                    "label": f"{parameter_raw} = {candidate_value}",
                    "parameter_key": parameter_key,
                    "execution_tier": _execution_tier_for_parameter(parameter_key),
                    "supported": supported,
                    "candidate_value": candidate_value,
                    "overrides": overrides,
                    "notes": notes,
                }
            )

    pack = {
        "version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "scenario_pack_id": "guadeloupe-sensitivity-default-v1",
        "territories": ["guadeloupe"],
        "source_workbook": str(workbook_path),
        "source_sheet": sheet_name,
        "workbook_total_runs": sheet_meta.get("workbook_total_runs"),
        "variables": variables,
        "scenarios": scenarios,
        "scenario_count": len(scenarios),
        "matches_workbook_total_runs": (
            sheet_meta.get("workbook_total_runs") is None
            or int(sheet_meta.get("workbook_total_runs")) == len(scenarios)
        ),
    }
    return raw_rows, pack


def write_outputs(output_dir: Path, raw_rows: list[dict[str, Any]], pack: dict[str, Any]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)

    csv_path = output_dir / RAW_CSV_NAME
    json_path = output_dir / RAW_JSON_NAME
    pack_path = output_dir / PACK_JSON_NAME

    if raw_rows:
        header = list(raw_rows[0].keys())
        with csv_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=header)
            writer.writeheader()
            for row in raw_rows:
                writer.writerow({key: _clean_text(value) for key, value in row.items()})

    json_path.write_text(
        json.dumps(
            {
                "generated_at": pack["generated_at"],
                "source_workbook": pack["source_workbook"],
                "source_sheet": pack["source_sheet"],
                "variables": pack["variables"],
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    pack_path.write_text(json.dumps(pack, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Export the sensitivity workbook sheet into versioned CSV/JSON files")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK_PATH)
    parser.add_argument("--sheet", type=str, default=DEFAULT_SHEET_NAME)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()

    raw_rows, pack = build_outputs(args.workbook, args.sheet)
    write_outputs(args.output_dir, raw_rows, pack)

    print(f"Workbook: {args.workbook}")
    print(f"Sheet: {args.sheet}")
    print(f"Variables exported: {len(pack['variables'])}")
    print(f"Scenarios exported: {len(pack['scenarios'])}")
    print(f"Output dir: {args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
